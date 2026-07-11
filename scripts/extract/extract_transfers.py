"""Collects and vendors the interchange transfer-times sheet.

Downloads the public Google Sheet, keeps the raw xlsx, and writes the Timing
sheet as CSV plus the remarks as text. See docs/SOURCES.md for provenance.

The CSV is normalized to current station codes; the xlsx preserves the
original. Transformations:
  - CE1/CE2 renamed to CC34/CC33 (Circle Line loop closure renumbering);
  - the two Tanah Merah "CG" rows dropped: 10 s cross-platform to the Changi
    branch, which shares stop EW4 in our feed, making the transfer implicit.
"""

import csv
from pathlib import Path

import openpyxl
import requests

SHEET_ID = "1e-Tuf6rHBFsgsuFN7XqbFL8ec_vdRjQw"
URL = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"
OUT = Path(__file__).parents[2] / "data" / "raw" / "transfers"

RENAMES = {"CE1": "CC34", "CE2": "CC33"}
DROP_CODES = {"CG"}
CODE_COLS = (2, 4)  # Start Code, End Code


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    xlsx = OUT / "transfer_times.xlsx"
    xlsx.write_bytes(requests.get(URL, timeout=60).content)

    wb = openpyxl.load_workbook(xlsx)

    kept = dropped = 0
    with open(OUT / "timing.csv", "w", newline="") as f:
        w = csv.writer(f)
        for row in wb["Timing"].iter_rows(values_only=True):
            if not any(v is not None for v in row):
                continue
            row = list(row)
            if any(row[i] in DROP_CODES for i in CODE_COLS):
                dropped += 1
                continue
            for i in CODE_COLS:
                row[i] = RENAMES.get(row[i], row[i])
            w.writerow(row)
            kept += 1
    print(f"kept {kept - 1} timing rows, dropped {dropped} (CG platform rows)")

    remarks = [
        str(row[0]).strip()
        for row in wb["Extra remarks"].iter_rows(values_only=True)
        if row[0]
    ]
    (OUT / "remarks.txt").write_text("\n".join(remarks) + "\n")
    print(f"{len(remarks)} remarks")


if __name__ == "__main__":
    main()
